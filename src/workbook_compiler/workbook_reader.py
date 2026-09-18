"""Workbook structure readers for the ECCS Workbook Compiler."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import openpyxl
from pyxlsb import open_workbook


def sha256_file(path: Path) -> str:
    """Return the SHA-256 hash of a file."""
    digest = hashlib.sha256()

    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def inspect_file_signature(path: Path) -> bytes:
    """Return the first eight bytes of the source file."""
    with path.open("rb") as stream:
        return stream.read(8)


def read_xlsb(path: Path) -> list[dict[str, Any]]:
    """Read worksheet metadata from an XLSB workbook."""
    sheets: list[dict] = []

    with open_workbook(path) as workbook:
        for name in workbook.sheets:
            sheets.append(
                {
                    "name": name,
                    "visibility": "unknown",
                    "classification": "unknown",
                }
            )

    return sheets


def read_xlsm_or_xlsx(path: Path) -> list[dict[str, Any]]:
    """Read worksheet metadata from an XLSM or XLSX workbook."""
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.lower() == ".xlsm",
    )

    sheets: list[dict] = []

    for worksheet in workbook.worksheets:
        sheets.append(
            {
                "name": worksheet.title,
                "visibility": "visible",
                "row_count": worksheet.max_row or 0,
                "column_count": worksheet.max_column or 0,
                "classification": "unknown",
            }
        )

    workbook.close()

    return sheets


def read_workbook(path: str) -> dict[str, Any]:
    """Validate and read a supported workbook."""
    source = Path(path)

    if not source.exists():
        raise FileNotFoundError(f"Workbook not found: {source}")

    extension = source.suffix.lower()

    if extension == ".xlsb":
        worksheets = read_xlsb(source)

    elif extension in {".xlsm", ".xlsx"}:
        worksheets = read_xlsm_or_xlsx(source)

    else:
        raise ValueError(
            f"Unsupported workbook type: {extension}"
        )

    return {
        "filename": source.name,
        "extension": extension,
        "sha256": sha256_file(source),
        "file_size_bytes": source.stat().st_size,
        "worksheets": worksheets,
    }
